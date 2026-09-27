import os
import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt

from app import app
from services.extractor import extract_text_from_file
from services.parser import parse_resume_text
from services.docx_builder import build_resume_docx

def test_pipeline():
    print("=== 1. CREATING SAMPLE PPTX RESUME ===")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6]) # blank layout
    
    # Title shape
    txBox = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "SARAH JENNINGS"
    
    p2 = tf.add_paragraph()
    p2.text = "Vice President of Cloud Operations"

    p3 = tf.add_paragraph()
    p3.text = "sarah.jennings@enterprise.io | +1 (555) 782-1209 | Dallas, TX | linkedin.com/in/sarahjennings"

    # Summary & Experience shape
    txBox2 = slide.shapes.add_textbox(Inches(1), Inches(1.8), Inches(8), Inches(4))
    tf2 = txBox2.text_frame
    
    p_sum_h = tf2.paragraphs[0]
    p_sum_h.text = "EXECUTIVE SUMMARY"
    
    p_sum = tf2.add_paragraph()
    p_sum.text = "Transformational technology executive with 15+ years delivering multi-million dollar cloud infrastructure and operational scalability."
    
    p_exp_h = tf2.add_paragraph()
    p_exp_h.text = "WORK EXPERIENCE"

    p_job1 = tf2.add_paragraph()
    p_job1.text = "VP of Infrastructure at CloudScale Inc (2020 - Present)"
    
    p_b1 = tf2.add_paragraph()
    p_b1.text = "• Scaled operations from 5 to 50 data centers globally with 99.999% uptime."
    
    p_b2 = tf2.add_paragraph()
    p_b2.text = "• Reduced cloud infrastructure spend by $3.2M annually through automated workload balancing."

    p_sk_h = tf2.add_paragraph()
    p_sk_h.text = "CORE SKILLS"
    p_sk = tf2.add_paragraph()
    p_sk.text = "Cloud Strategy, DevOps, Microservices, Infrastructure as Code, Kubernetes, Budget Management"

    pptx_path = "C:/Users/ITkey/.gemini/antigravity/scratch/resume_converter_app/uploads/test_sarah.pptx"
    prs.save(pptx_path)
    print(f"Sample PPTX saved to {pptx_path}")

    print("\n=== 2. EXTRACTING TEXT FROM PPTX ===")
    extracted = extract_text_from_file(pptx_path)
    assert extracted["success"] is True, f"Extraction failed: {extracted.get('error')}"
    print(f"Extracted raw text length: {len(extracted['raw_text'])} characters")

    print("\n=== 3. PARSING EXTRACTED TEXT ===")
    parsed = parse_resume_text(extracted["raw_text"])
    print(f"Parsed Name: {parsed['full_name']}")
    print(f"Parsed Title: {parsed['target_title']}")
    print(f"Parsed Email: {parsed['contact']['email']}")
    print(f"Parsed Phone: {parsed['contact']['phone']}")
    print(f"Parsed Skills Count: {len(parsed['skills'])}")
    print(f"Parsed Experience Count: {len(parsed['experience'])}")
    assert "SARAH" in parsed["full_name"].upper(), "Failed to extract candidate name"
    assert len(parsed["skills"]) > 0, "Failed to parse skills"
    assert parsed["contact"]["email"] == "sarah.jennings@enterprise.io", "Failed to extract candidate email"

    print("\n=== 4. GENERATING WORD DOCUMENT WITH LOGO IN HEADER ===")
    logo_file = "C:/Users/ITkey/.gemini/antigravity/scratch/resume_converter_app/static/images/company_logo.png"
    out_docx = "C:/Users/ITkey/.gemini/antigravity/scratch/resume_converter_app/output/Sarah_Jennings_Converted.docx"
    doc_res = build_resume_docx(
        resume_data=parsed,
        output_path=out_docx,
        logo_path=logo_file,
        header_alignment="dual",
        header_tagline="CONFIDENTIAL CANDIDATE DOSSIER | KEY DYNAMICS SOLUTIONS",
        palette_name="key_dynamics"
    )
    assert os.path.exists(doc_res), "Generated docx does not exist"
    print(f"Word document generated: {doc_res} ({os.path.getsize(doc_res)} bytes)")

    print("\n=== 5. TESTING FLASK API ENDPOINTS ===")
    client = app.test_client()
    
    # Test GET /
    res_index = client.get("/")
    assert res_index.status_code == 200
    assert b"Key Dynamics Solutions" in res_index.data
    print("GET / index page: 200 OK")

    # Test POST /api/upload
    with open(pptx_path, "rb") as f:
        res_upload = client.post("/api/upload", data={"file": (f, "test_sarah.pptx")}, content_type="multipart/form-data")
    assert res_upload.status_code == 200
    upload_json = res_upload.get_json()
    assert upload_json["success"] is True
    print("POST /api/upload PPTX: 200 OK")

    # Test POST /api/generate-docx
    gen_payload = {
        "resume_data": upload_json["structured_data"],
        "header_alignment": "right",
        "header_tagline": "KEY DYNAMICS SOLUTIONS - TRANSFORM YOUR OPERATION",
        "palette_name": "key_dynamics",
        "logo_width": 2.4
    }
    res_gen = client.post("/api/generate-docx", json=gen_payload)
    assert res_gen.status_code == 200
    gen_json = res_gen.get_json()
    assert gen_json["success"] is True
    assert "download_url" in gen_json
    print(f"POST /api/generate-docx: 200 OK -> {gen_json['download_url']}")

    # Test GET download
    res_download = client.get(gen_json["download_url"])
    assert res_download.status_code == 200
    assert res_download.content_length > 10000
    print(f"GET {gen_json['download_url']}: 200 OK ({res_download.content_length} bytes)")

    print("\nALL AUTOMATED TESTS PASSED SUCCESSFULLY!")

if __name__ == "__main__":
    test_pipeline()
