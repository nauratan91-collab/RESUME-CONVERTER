import os
from pathlib import Path
from typing import Dict, Any, Optional

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from services.parser import classify_title_and_company

# Color Schemes
COLOR_PALETTES = {
    "key_dynamics": {
        "primary": RGBColor(24, 34, 124),     # Deep Navy
        "primary_hex": "18227C",
        "accent": RGBColor(0, 181, 184),      # Bright Cyan / Teal
        "accent_hex": "00B5B8",
        "secondary": RGBColor(71, 85, 105),   # Slate Grey
        "secondary_hex": "475569",
        "text": RGBColor(30, 41, 59),         # Charcoal text
        "light_bg": "F0FDF4"
    },
    "modern_teal": {
        "primary": RGBColor(15, 118, 110),
        "primary_hex": "0F766E",
        "accent": RGBColor(6, 182, 212),
        "accent_hex": "06B6D4",
        "secondary": RGBColor(51, 65, 85),
        "secondary_hex": "334155",
        "text": RGBColor(30, 41, 59),
        "light_bg": "F0FDFA"
    },
    "classic_blue": {
        "primary": RGBColor(30, 58, 138),
        "primary_hex": "1E3A8A",
        "accent": RGBColor(59, 130, 246),
        "accent_hex": "3B82F6",
        "secondary": RGBColor(75, 85, 99),
        "secondary_hex": "4B5563",
        "text": RGBColor(17, 24, 39),
        "light_bg": "F8FAFC"
    },
    "sleek_dark": {
        "primary": RGBColor(17, 24, 39),
        "primary_hex": "111827",
        "accent": RGBColor(100, 116, 139),
        "accent_hex": "64748B",
        "secondary": RGBColor(75, 85, 99),
        "secondary_hex": "4B5563",
        "text": RGBColor(31, 41, 55),
        "light_bg": "F9FAFB"
    }
}


def add_bottom_border_to_paragraph(p, hex_color="00B5B8", sz="12"):
    """Adds a stylish bottom border line to a paragraph (like a section divider)"""
    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{sz}" w:space="4" w:color="{hex_color}"/></w:pBdr>')
    pPr.append(pBdr)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets padding/margins for table cells in dxa (1 pt = 20 dxa)"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)


def build_resume_docx(
    resume_data: Dict[str, Any],
    output_path: str,
    logo_path: Optional[str] = None,
    header_alignment: str = "right",      # 'right', 'left', 'center', 'dual'
    header_tagline: str = "CONFIDENTIAL CANDIDATE DOSSIER | KEY DYNAMICS SOLUTIONS",
    palette_name: str = "key_dynamics",
    logo_width_inches: float = 2.3,
    include_footer: bool = True
) -> str:
    """
    Builds a professional Word (.docx) document formatted with:
    - Company logo in header
    - Customizable header alignment & optional tagline
    - Executive typography and branded color scheme
    - Candidate info, summary, experience, education, skills, certifications
    """
    palette = COLOR_PALETTES.get(palette_name, COLOR_PALETTES["key_dynamics"])
    doc = docx.Document()

    # Configure Margins
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    # 1. SETUP HEADER WITH COMPANY LOGO
    header = section.header
    
    # Check if logo exists
    has_logo = logo_path and os.path.exists(logo_path)

    if has_logo:
        if header_alignment == "dual":
            # 2-column table: Logo on left, Tagline/Company info on right
            tbl = header.add_table(rows=1, cols=2, width=Inches(6.8))
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            
            # Left cell for logo
            cell_left = tbl.cell(0, 0)
            cell_left.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p_logo = cell_left.paragraphs[0]
            p_logo.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_logo.paragraph_format.space_after = Pt(0)
            run_logo = p_logo.add_run()
            run_logo.add_picture(logo_path, width=Inches(logo_width_inches))

            # Right cell for tagline
            cell_right = tbl.cell(0, 1)
            cell_right.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p_tag = cell_right.paragraphs[0]
            p_tag.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_tag.paragraph_format.space_after = Pt(0)
            if header_tagline:
                r_tag = p_tag.add_run(header_tagline)
                r_tag.font.size = Pt(8)
                r_tag.font.name = "Calibri"
                r_tag.font.color.rgb = palette["secondary"]
                r_tag.font.bold = True

            # Add separator border under header
            p_sep = header.add_paragraph()
            p_sep.paragraph_format.space_before = Pt(4)
            p_sep.paragraph_format.space_after = Pt(0)
            add_bottom_border_to_paragraph(p_sep, hex_color=palette["accent_hex"], sz="8")

        elif header_alignment == "center":
            p = header.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run()
            run.add_picture(logo_path, width=Inches(logo_width_inches))
            
            if header_tagline:
                p_tag = header.add_paragraph()
                p_tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_tag.paragraph_format.space_after = Pt(4)
                r_tag = p_tag.add_run(header_tagline)
                r_tag.font.size = Pt(8)
                r_tag.font.color.rgb = palette["secondary"]
                r_tag.font.bold = True
                add_bottom_border_to_paragraph(p_tag, hex_color=palette["accent_hex"], sz="8")
            else:
                add_bottom_border_to_paragraph(p, hex_color=palette["accent_hex"], sz="8")

        elif header_alignment == "left":
            p = header.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run()
            run.add_picture(logo_path, width=Inches(logo_width_inches))
            
            if header_tagline:
                p_tag = header.add_paragraph()
                p_tag.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p_tag.paragraph_format.space_after = Pt(4)
                r_tag = p_tag.add_run(header_tagline)
                r_tag.font.size = Pt(8)
                r_tag.font.color.rgb = palette["secondary"]
                r_tag.font.bold = True
                add_bottom_border_to_paragraph(p_tag, hex_color=palette["accent_hex"], sz="8")
            else:
                add_bottom_border_to_paragraph(p, hex_color=palette["accent_hex"], sz="8")

        else: # Default: Right-aligned logo
            p = header.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run()
            run.add_picture(logo_path, width=Inches(logo_width_inches))
            
            p_sep = header.add_paragraph()
            p_sep.paragraph_format.space_before = Pt(2)
            p_sep.paragraph_format.space_after = Pt(0)
            if header_tagline:
                p_sep.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                r_tag = p_sep.add_run(header_tagline)
                r_tag.font.size = Pt(7.5)
                r_tag.font.color.rgb = palette["secondary"]
                r_tag.font.bold = True
            add_bottom_border_to_paragraph(p_sep, hex_color=palette["accent_hex"], sz="8")

    # 2. SETUP FOOTER
    if include_footer:
        footer = section.footer
        p_ft = footer.paragraphs[0]
        p_ft.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_ft.paragraph_format.space_before = Pt(6)
        p_ft.paragraph_format.space_after = Pt(0)
        
        # Add subtle top border to footer
        pPr = p_ft._p.get_or_add_pPr()
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="3" w:color="{palette["accent_hex"]}"/></w:pBdr>')
        pPr.append(pBdr)

        r_ft = p_ft.add_run("Key Dynamics Solutions  •  Transform Your Operation  •  Confidential")
        r_ft.font.size = Pt(8)
        r_ft.font.name = "Calibri"
        r_ft.font.color.rgb = palette["secondary"]

    # 3. CANDIDATE HEADER / NAME
    p_name = doc.add_paragraph()
    p_name.paragraph_format.space_before = Pt(4)
    p_name.paragraph_format.space_after = Pt(2)
    
    candidate_name = resume_data.get("full_name", "CANDIDATE NAME").upper()
    r_name = p_name.add_run(candidate_name)
    r_name.font.size = Pt(22)
    r_name.font.name = "Calibri"
    r_name.font.bold = True
    r_name.font.color.rgb = palette["primary"]

    # Candidate Target Title
    target_title = resume_data.get("target_title", "").strip()
    if target_title:
        p_title = doc.add_paragraph()
        p_title.paragraph_format.space_before = Pt(0)
        p_title.paragraph_format.space_after = Pt(4)
        r_title = p_title.add_run(target_title)
        r_title.font.size = Pt(13)
        r_title.font.name = "Calibri"
        r_title.font.bold = True
        r_title.font.color.rgb = palette["accent"]

    # Contact Info Bar
    contact = resume_data.get("contact", {})
    contact_parts = []
    if contact.get("email"):
        contact_parts.append(contact["email"])
    if contact.get("phone"):
        contact_parts.append(contact["phone"])
    if contact.get("location"):
        contact_parts.append(contact["location"])
    if contact.get("linkedin"):
        contact_parts.append(contact["linkedin"])
    if contact.get("website"):
        contact_parts.append(contact["website"])

    if contact_parts:
        p_contact = doc.add_paragraph()
        p_contact.paragraph_format.space_before = Pt(0)
        p_contact.paragraph_format.space_after = Pt(10)
        
        contact_text = "   |   ".join(contact_parts)
        r_contact = p_contact.add_run(contact_text)
        r_contact.font.size = Pt(9.5)
        r_contact.font.name = "Calibri"
        r_contact.font.color.rgb = palette["secondary"]

    # Helper function for section headings
    def add_section_header(title: str):
        p_sec = doc.add_paragraph()
        p_sec.paragraph_format.space_before = Pt(14)
        p_sec.paragraph_format.space_after = Pt(6)
        p_sec.paragraph_format.keep_with_next = True
        
        r_sec = p_sec.add_run(title.upper())
        r_sec.font.size = Pt(12)
        r_sec.font.name = "Calibri"
        r_sec.font.bold = True
        r_sec.font.color.rgb = palette["primary"]

        # Add section bottom border
        add_bottom_border_to_paragraph(p_sec, hex_color=palette["primary_hex"], sz="12")

    # 4. PROFESSIONAL SUMMARY
    summary_text = resume_data.get("summary", "").strip()
    if summary_text:
        add_section_header("Professional Summary")
        p_sum = doc.add_paragraph()
        p_sum.paragraph_format.space_before = Pt(4)
        p_sum.paragraph_format.space_after = Pt(8)
        p_sum.paragraph_format.line_spacing = 1.15
        r_sum = p_sum.add_run(summary_text)
        r_sum.font.size = Pt(10.5)
        r_sum.font.name = "Calibri"
        r_sum.font.color.rgb = palette["text"]

    # 5. CORE SKILLS
    skills = resume_data.get("skills", [])
    if skills:
        add_section_header("Core Competencies & Skills")
        
        # Format skills in a clean 2-column or 3-column table
        cols_count = 3 if len(skills) >= 6 else 2
        chunk_size = (len(skills) + cols_count - 1) // cols_count
        
        tbl_skills = doc.add_table(rows=chunk_size, cols=cols_count)
        tbl_skills.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        for idx, skill in enumerate(skills):
            r_idx = idx % chunk_size
            c_idx = idx // chunk_size
            if c_idx < cols_count:
                cell = tbl_skills.cell(r_idx, c_idx)
                set_cell_margins(cell, top=40, bottom=40, left=80, right=80)
                p_sk = cell.paragraphs[0]
                p_sk.paragraph_format.space_before = Pt(1)
                p_sk.paragraph_format.space_after = Pt(1)
                
                # Bullet symbol in accent color
                r_bullet = p_sk.add_run("▪  ")
                r_bullet.font.size = Pt(9.5)
                r_bullet.font.color.rgb = palette["accent"]
                
                r_text = p_sk.add_run(skill)
                r_text.font.size = Pt(10)
                r_text.font.name = "Calibri"
                r_text.font.color.rgb = palette["text"]

        p_spacer = doc.add_paragraph()
        p_spacer.paragraph_format.space_before = Pt(2)
        p_spacer.paragraph_format.space_after = Pt(2)

    # 6. PROFESSIONAL EXPERIENCE
    experiences = resume_data.get("experience", [])
    if experiences:
        add_section_header("Professional Experience")
        
        for exp in experiences:
            raw_title = exp.get("title", "Role").strip()
            raw_company = exp.get("company", "").strip()
            title, company = classify_title_and_company(raw_title, raw_company)
            dates = exp.get("dates", "").strip()
            location = exp.get("location", "").strip()
            bullets = exp.get("bullets", [])

            # Role & Dates line using a borderless 2-column table for alignment
            tbl_job = doc.add_table(rows=1, cols=2)
            tbl_job.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl_job.autofit = False

            # Cell Left: Role & Company
            cell_left = tbl_job.cell(0, 0)
            cell_left.width = Inches(4.8)
            set_cell_margins(cell_left, top=40, bottom=20, left=0, right=50)
            p_left = cell_left.paragraphs[0]
            p_left.paragraph_format.space_before = Pt(3)
            p_left.paragraph_format.space_after = Pt(1)
            
            r_role = p_left.add_run(title)
            r_role.font.size = Pt(11)
            r_role.font.name = "Calibri"
            r_role.font.bold = True
            r_role.font.color.rgb = palette["primary"]

            if company:
                r_comp = p_left.add_run(f"  |  {company}")
                r_comp.font.size = Pt(10.5)
                r_comp.font.name = "Calibri"
                r_comp.font.bold = True
                r_comp.font.color.rgb = palette["secondary"]

            # Cell Right: Dates & Location
            cell_right = tbl_job.cell(0, 1)
            cell_right.width = Inches(2.0)
            set_cell_margins(cell_right, top=40, bottom=20, left=50, right=0)
            p_right = cell_right.paragraphs[0]
            p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_right.paragraph_format.space_before = Pt(3)
            p_right.paragraph_format.space_after = Pt(1)

            date_str = dates
            if location:
                date_str = f"{dates}  •  {location}" if dates else location

            r_dates = p_right.add_run(date_str)
            r_dates.font.size = Pt(9.5)
            r_dates.font.name = "Calibri"
            r_dates.font.italic = True
            r_dates.font.color.rgb = palette["secondary"]

            # Bullets
            for bullet in bullets:
                if not bullet.strip():
                    continue
                p_b = doc.add_paragraph()
                p_b.paragraph_format.left_indent = Inches(0.25)
                p_b.paragraph_format.space_before = Pt(1)
                p_b.paragraph_format.space_after = Pt(2.5)
                p_b.paragraph_format.line_spacing = 1.15

                # Custom styled bullet symbol
                r_sym = p_b.add_run("•  ")
                r_sym.font.name = "Calibri"
                r_sym.font.bold = True
                r_sym.font.color.rgb = palette["accent"]

                r_btxt = p_b.add_run(bullet.strip())
                r_btxt.font.size = Pt(10)
                r_btxt.font.name = "Calibri"
                r_btxt.font.color.rgb = palette["text"]

            # Small space after job
            p_gap = doc.add_paragraph()
            p_gap.paragraph_format.space_before = Pt(2)
            p_gap.paragraph_format.space_after = Pt(2)

    # 7. EDUCATION
    educations = resume_data.get("education", [])
    if educations:
        add_section_header("Education & Credentials")
        
        for edu in educations:
            deg = edu.get("degree", "").strip()
            inst = edu.get("institution", "").strip()
            dates = edu.get("dates", "").strip()
            details = edu.get("details", "").strip()

            tbl_edu = doc.add_table(rows=1, cols=2)
            tbl_edu.alignment = WD_TABLE_ALIGNMENT.CENTER

            cell_l = tbl_edu.cell(0, 0)
            cell_l.width = Inches(4.8)
            set_cell_margins(cell_l, top=30, bottom=20, left=0, right=50)
            p_l = cell_l.paragraphs[0]
            p_l.paragraph_format.space_before = Pt(2)
            p_l.paragraph_format.space_after = Pt(1)

            r_deg = p_l.add_run(deg)
            r_deg.font.size = Pt(10.5)
            r_deg.font.name = "Calibri"
            r_deg.font.bold = True
            r_deg.font.color.rgb = palette["primary"]

            if inst:
                r_inst = p_l.add_run(f" — {inst}")
                r_inst.font.size = Pt(10)
                r_inst.font.name = "Calibri"
                r_inst.font.color.rgb = palette["secondary"]

            cell_r = tbl_edu.cell(0, 1)
            cell_r.width = Inches(2.0)
            set_cell_margins(cell_r, top=30, bottom=20, left=50, right=0)
            p_r = cell_r.paragraphs[0]
            p_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_r.paragraph_format.space_before = Pt(2)
            p_r.paragraph_format.space_after = Pt(1)

            if dates:
                r_d = p_r.add_run(dates)
                r_d.font.size = Pt(9.5)
                r_d.font.name = "Calibri"
                r_d.font.italic = True
                r_d.font.color.rgb = palette["secondary"]

            if details:
                p_det = doc.add_paragraph()
                p_det.paragraph_format.left_indent = Inches(0.2)
                p_det.paragraph_format.space_before = Pt(0)
                p_det.paragraph_format.space_after = Pt(3)
                r_det = p_det.add_run(details)
                r_det.font.size = Pt(9.5)
                r_det.font.name = "Calibri"
                r_det.font.color.rgb = palette["text"]

    # 8. CERTIFICATIONS / AWARDS
    certifications = resume_data.get("certifications", [])
    if certifications:
        add_section_header("Certifications & Honors")
        for cert in certifications:
            p_c = doc.add_paragraph()
            p_c.paragraph_format.left_indent = Inches(0.2)
            p_c.paragraph_format.space_before = Pt(1)
            p_c.paragraph_format.space_after = Pt(2)

            r_tick = p_c.add_run("✔  ")
            r_tick.font.color.rgb = palette["accent"]
            r_tick.font.size = Pt(9.5)

            r_ctxt = p_c.add_run(cert)
            r_ctxt.font.size = Pt(10)
            r_ctxt.font.name = "Calibri"
            r_ctxt.font.color.rgb = palette["text"]

    # 9. PROJECTS
    projects = resume_data.get("projects", [])
    if projects:
        add_section_header("Key Projects")
        for proj in projects:
            p_pj = doc.add_paragraph()
            p_pj.paragraph_format.space_before = Pt(3)
            p_pj.paragraph_format.space_after = Pt(1)

            r_pjn = p_pj.add_run(proj.get("name", "Project"))
            r_pjn.font.size = Pt(10.5)
            r_pjn.font.name = "Calibri"
            r_pjn.font.bold = True
            r_pjn.font.color.rgb = palette["primary"]

            if proj.get("description"):
                p_pjd = doc.add_paragraph()
                p_pjd.paragraph_format.left_indent = Inches(0.2)
                p_pjd.paragraph_format.space_before = Pt(0)
                p_pjd.paragraph_format.space_after = Pt(2)
                r_pdesc = p_pjd.add_run(proj["description"])
                r_pdesc.font.size = Pt(9.5)
                r_pdesc.font.name = "Calibri"
                r_pdesc.font.color.rgb = palette["text"]

            for b in proj.get("bullets", []):
                p_pjb = doc.add_paragraph()
                p_pjb.paragraph_format.left_indent = Inches(0.25)
                p_pjb.paragraph_format.space_before = Pt(1)
                p_pjb.paragraph_format.space_after = Pt(2)
                r_b = p_pjb.add_run("•  ")
                r_b.font.color.rgb = palette["accent"]
                r_btxt = p_pjb.add_run(b)
                r_btxt.font.size = Pt(9.5)
                r_btxt.font.name = "Calibri"
                r_btxt.font.color.rgb = palette["text"]

    # Save generated document
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path
