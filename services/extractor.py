import os
from pathlib import Path
from pypdf import PdfReader
from pptx import Presentation
from docx import Document
from PIL import Image

def extract_text_from_file(file_path: str, gemini_api_key: str = None) -> dict:
    """
    Extract text and structured details from PDF, PPTX, DOCX, TXT, or Image files.
    Returns:
        {
            "success": bool,
            "filename": str,
            "file_type": str,
            "raw_text": str,
            "page_count": int,
            "error": str (optional),
            "ai_parsed": bool (optional),
            "structured_data": dict (optional if AI directly parsed it)
        }
    """
    path = Path(file_path)
    ext = path.suffix.lower()
    filename = path.name

    if not path.exists():
        return {"success": False, "filename": filename, "file_type": ext, "error": f"File not found: {file_path}"}

    try:
        if ext == ".pdf":
            return extract_from_pdf(file_path, gemini_api_key)
        elif ext in [".pptx", ".ppt"]:
            return extract_from_pptx(file_path)
        elif ext in [".docx", ".doc"]:
            return extract_from_docx(file_path)
        elif ext in [".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff"]:
            return extract_from_image(file_path, gemini_api_key)
        elif ext in [".txt", ".rtf", ".md"]:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
            return {
                "success": True,
                "filename": filename,
                "file_type": ext,
                "raw_text": content,
                "page_count": 1
            }
        else:
            return {"success": False, "filename": filename, "file_type": ext, "error": f"Unsupported format: {ext}"}
    except Exception as e:
        return {"success": False, "filename": filename, "file_type": ext, "error": str(e)}


def extract_from_pdf(file_path: str, gemini_api_key: str = None) -> dict:
    reader = PdfReader(file_path)
    pages_text = []
    for idx, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        pages_text.append(text)
    
    combined_text = "\n\n".join(pages_text).strip()

    # If the PDF is scanned or has almost no extracted text, and AI key is provided, try Gemini Vision
    if len(combined_text) < 100 and gemini_api_key:
        from services.ai_enhancer import parse_document_with_ai
        ai_res = parse_document_with_ai(file_path, gemini_api_key, mime_type="application/pdf")
        if ai_res.get("success"):
            return ai_res

    return {
        "success": True,
        "filename": Path(file_path).name,
        "file_type": ".pdf",
        "raw_text": combined_text,
        "page_count": len(reader.pages)
    }


def extract_from_pptx(file_path: str) -> dict:
    prs = Presentation(file_path)
    slides_text = []
    
    for slide_idx, slide in enumerate(prs.slides):
        slide_lines = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = paragraph.text.strip()
                    if line:
                        slide_lines.append(line)
            elif shape.has_table:
                for row in shape.table.rows:
                    row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_texts:
                        slide_lines.append(" | ".join(row_texts))
        
        if slide_lines:
            slides_text.append("\n".join(slide_lines))

    combined_text = "\n\n--- Slide ---\n\n".join(slides_text)

    return {
        "success": True,
        "filename": Path(file_path).name,
        "file_type": ".pptx",
        "raw_text": combined_text,
        "page_count": len(prs.slides)
    }


def extract_from_docx(file_path: str) -> dict:
    doc = Document(file_path)
    paragraphs = []
    
    for p in doc.paragraphs:
        if p.text.strip():
            paragraphs.append(p.text.strip())

    for table in doc.tables:
        for row in table.rows:
            row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_texts:
                paragraphs.append(" | ".join(row_texts))

    combined_text = "\n".join(paragraphs)

    return {
        "success": True,
        "filename": Path(file_path).name,
        "file_type": ".docx",
        "raw_text": combined_text,
        "page_count": 1
    }


def extract_from_image(file_path: str, gemini_api_key: str = None) -> dict:
    filename = Path(file_path).name
    # If Gemini API key is present, use Multimodal Vision
    if gemini_api_key:
        from services.ai_enhancer import parse_document_with_ai
        img_type = "image/png" if file_path.lower().endswith(".png") else "image/jpeg"
        return parse_document_with_ai(file_path, gemini_api_key, mime_type=img_type)

    # If no API key is provided, return guidance or inspect image details
    try:
        with Image.open(file_path) as img:
            w, h = img.size
            fmt = img.format
    except Exception as e:
        return {"success": False, "filename": filename, "file_type": "image", "error": f"Invalid image: {str(e)}"}

    return {
        "success": True,
        "filename": filename,
        "file_type": "image",
        "raw_text": f"[Image Resume: {filename} ({w}x{h} px)]\nTip: For scanned or graphic image resumes, enter your Gemini API key in Settings for AI OCR parsing, or you can paste / fill in the extracted details below.",
        "page_count": 1,
        "requires_ocr": True
    }
